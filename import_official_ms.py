import os
import sys
import json
import shutil
from datetime import datetime
import openpyxl
import pandas as pd
import requests

sys.stdout.reconfigure(encoding='utf-8')

EXCEL_SOURCE = r"C:\Users\peera\Desktop\AntigravityProject\New folder\ประกาศรายชื่อนักเรียนที่มีเวลาเรียนไม่ถึง 80%-1_2569.xlsx"
SHEET_NAME = "M.S. Report (อัปเดตแล้ว)"

TEACHERHUB_DIR = r"c:\Users\peera\Desktop\AntigravityProject\TeacherHub-Backend"
ASSESSMENT_BACKEND_DIR = r"c:\Users\peera\Desktop\AntigravityProject\AssessmentHub\SgsNextschool\backend"

WP16_JSON = os.path.join(TEACHERHUB_DIR, "wp16_pending_tasks.json")
WP16_EXCEL_BACKEND = os.path.join(TEACHERHUB_DIR, "WP16_งานค้าง_ต้นทาง.xlsx")
WP16_EXCEL_DESKTOP = os.path.join(os.environ.get('USERPROFILE', r'C:\Users\peera'), 'Desktop', 'WP16_งานค้าง_ต้นทาง.xlsx')

OFFICIAL_MS_JSON_TEACHERHUB = os.path.join(TEACHERHUB_DIR, "sgs_module", "official_ms_list.json")
OFFICIAL_MS_JSON_ASSESSMENT = os.path.join(ASSESSMENT_BACKEND_DIR, "official_ms_list.json")

GAS_URL = "https://script.google.com/macros/s/AKfycbxzpP9b_eBJUU5KaNX1CbMOLHygMsrUdO7earro-bQIs8lMS9H6YM6Z6mlamm3jJd1fDQ/exec"

def run_import():
    print(f"📖 กำลังอ่านข้อมูลจาก: {EXCEL_SOURCE} (แท็บ: {SHEET_NAME})...")
    if not os.path.exists(EXCEL_SOURCE):
        print(f"❌ ไม่พบไฟล์: {EXCEL_SOURCE}")
        return False

    wb = openpyxl.load_workbook(EXCEL_SOURCE, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"❌ ไม่พบแท็บ: {SHEET_NAME} ในไฟล์")
        return False

    ws = wb[SHEET_NAME]

    ms_records = {}
    allowed_exemptions = {}
    by_subject = {}

    for r in range(2, ws.max_row + 1):
        sid = str(ws.cell(r, 2).value or '').strip()
        sname = str(ws.cell(r, 3).value or '').strip()
        clevel = str(ws.cell(r, 4).value or '').strip()
        scode = str(ws.cell(r, 6).value or '').strip()
        subname = str(ws.cell(r, 7).value or '').strip()
        teacher = str(ws.cell(r, 13).value or '').strip()
        remark = str(ws.cell(r, 14).value or '').strip()
        allowed = ws.cell(r, 15).value

        if not sid or not scode:
            continue

        special_id = f"{scode}{sid}"

        record = {
            "special_id": special_id,
            "academic_year": "2569",
            "semester": "1",
            "subject_code": scode,
            "subject_name": subname,
            "teacher_name": teacher,
            "class_level": clevel,
            "student_id": sid,
            "student_name": sname,
            "old_score": "",
            "old_grade": "มส",
            "pending_task": remark,
            "remark": "[มส. ประกาศทางการ]",
            "is_manual": False,
            "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }

        if allowed is True:
            allowed_exemptions[special_id] = {
                **record,
                "allowed": True,
                "pending_task": f"{remark} [อนุญาตให้สอบ]",
                "remark": "[อนุญาตให้เข้าสอบ]"
            }
        else:
            ms_records[special_id] = record
            if scode not in by_subject:
                by_subject[scode] = []
            by_subject[scode].append(sid)

    print(f"✅ ประมวลผลเสร็จสิ้น:")
    print(f"   - รายชื่อ มส. ทางการ (ไม่อนุญาตให้สอบ): {len(ms_records)} รายการ")
    print(f"   - รายชื่อที่ได้รับอนุญาตให้สอบ (ยกเว้น มส.): {len(allowed_exemptions)} รายการ")
    print(f"   - จำนวนวิชาทั้งหมดที่พบ: {len(by_subject)} วิชา")

    # 1. บันทึก official_ms_list.json สำหรับ validator ใช้ตรวจสอบ
    official_data = {
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source": os.path.basename(EXCEL_SOURCE),
        "sheet": SHEET_NAME,
        "total_ms": len(ms_records),
        "total_allowed": len(allowed_exemptions),
        "ms_records": ms_records,
        "allowed_exemptions": allowed_exemptions,
        "by_subject": by_subject
    }

    os.makedirs(os.path.dirname(OFFICIAL_MS_JSON_TEACHERHUB), exist_ok=True)
    with open(OFFICIAL_MS_JSON_TEACHERHUB, "w", encoding="utf-8") as f:
        json.dump(official_data, f, ensure_ascii=False, indent=2)
    print(f"💾 บันทึก official_ms_list.json ที่: {OFFICIAL_MS_JSON_TEACHERHUB}")

    os.makedirs(os.path.dirname(OFFICIAL_MS_JSON_ASSESSMENT), exist_ok=True)
    with open(OFFICIAL_MS_JSON_ASSESSMENT, "w", encoding="utf-8") as f:
        json.dump(official_data, f, ensure_ascii=False, indent=2)
    print(f"💾 บันทึก official_ms_list.json ที่: {OFFICIAL_MS_JSON_ASSESSMENT}")

    # 2. บันทึกลงใน wp16_pending_tasks.json
    # อ่าน tasks เดิม (ถ้ามีรายการของผู้ใช้ที่ครูกรอกเพิ่มเอง ให้คงไว้)
    existing_db = {}
    if os.path.exists(WP16_JSON):
        try:
            with open(WP16_JSON, "r", encoding="utf-8") as f:
                existing_db = json.load(f)
        except Exception:
            existing_db = {}

    # รวมข้อมูล: ใส่ มส. ทางการเข้าไป (ถ้ามี task เดิมอยู่แล้วและเป็น manual ให้คง task เดิมไว้ แต่ปรับเกรดเป็น มส)
    merged_db = dict(existing_db)
    # ลบข้อมูลทดสอบ เช่น 54545
    if "ส3120554545" in merged_db and merged_db["ส3120554545"].get("student_name") == "6565":
        del merged_db["ส3120554545"]

    for sp_id, item in ms_records.items():
        if sp_id in merged_db and merged_db[sp_id].get("is_manual"):
            # คงค่างานเดิมที่ครูอาจกรอกไว้
            prev = merged_db[sp_id]
            merged_db[sp_id] = {
                **item,
                "pending_task": prev.get("pending_task") or item["pending_task"],
                "old_score": prev.get("old_score") or item["old_score"]
            }
        else:
            merged_db[sp_id] = item

    with open(WP16_JSON, "w", encoding="utf-8") as f:
        json.dump(merged_db, f, ensure_ascii=False, indent=2)
    print(f"💾 อัปเดต {WP16_JSON} เรียบร้อย (รวม {len(merged_db)} รายการ)")

    # 3. บันทึกไฟล์ Excel WP16_งานค้าง_ต้นทาง.xlsx (14 คอลัมน์ ตรงตามหัวตารางมาตรฐาน)
    thai_col_map = {
        'special_id': 'เลขเฉพาะ',
        'academic_year': 'ปีการศึกษา',
        'semester': 'ภาคเรียน',
        'subject_code': 'รหัสวิชา',
        'subject_name': 'ชื่อวิชา',
        'teacher_name': 'ครูผู้สอน',
        'class_level': 'ชั้น/ห้อง',
        'student_id': 'เลขประจำตัว',
        'student_name': 'ชื่อ-นามสกุล',
        'old_score': 'คะแนนเดิม',
        'old_grade': 'ผลการเรียนเดิม',
        'pending_task': 'งานค้าง',
        'remark': 'หมายเหตุ',
        'updated_at': 'วันที่บันทึก'
    }
    preferred_cols = list(thai_col_map.keys())
    rows = []
    for k, v in merged_db.items():
        r = dict(v)
        rows.append(r)
    df = pd.DataFrame(rows)
    for col in preferred_cols:
        if col not in df.columns:
            df[col] = ''
    df = df[preferred_cols]
    df = df.rename(columns=thai_col_map)

    with pd.ExcelWriter(WP16_EXCEL_BACKEND, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='WP16_งานค้าง', index=False)
    print(f"💾 บันทึก Excel: {WP16_EXCEL_BACKEND}")

    try:
        shutil.copy2(WP16_EXCEL_BACKEND, WP16_EXCEL_DESKTOP)
        print(f"💾 สำเนาไฟล์ไปยัง Desktop: {WP16_EXCEL_DESKTOP}")
    except Exception as e:
        print(f"⚠️ ไม่สามารถคัดลอกไปยัง Desktop ได้: {e}")

    # 4. ซิงค์ขึ้น Google Sheet Sgsnextschool แท็บ WP16_งานค้าง
    print(f"\n🚀 กำลังซิงค์ข้อมูล {len(merged_db)} รายการขึ้น Google Sheet Sgsnextschool...")
    items_to_sync = []
    for k, v in merged_db.items():
        yr = str(v.get('academic_year', '2569'))
        sem = str(v.get('semester', '1'))
        items_to_sync.append({
            'special_id': v.get('special_id') or k,
            'specialId': v.get('special_id') or k,
            'academic_year': yr,
            'semester': sem,
            'subject_code': v.get('subject_code', ''),
            'subjCode': v.get('subject_code', ''),
            'subject_name': v.get('subject_name', ''),
            'subjName': v.get('subject_name', ''),
            'teacher_name': v.get('teacher_name', ''),
            'teacherName': v.get('teacher_name', ''),
            'class_level': v.get('class_level', ''),
            'classLevel': v.get('class_level', ''),
            'student_id': v.get('student_id', ''),
            'stuId': v.get('student_id', ''),
            'student_name': v.get('student_name', ''),
            'stuName': v.get('student_name', ''),
            'old_score': v.get('old_score', ''),
            'old_grade': v.get('old_grade', 'มส'),
            'oldGrade': v.get('old_grade', 'มส'),
            'pending_task': v.get('pending_task', ''),
            'pendingTask': v.get('pending_task', ''),
            'remark': v.get('remark', ''),
            'updated_at': v.get('updated_at', '')
        })

    payload = json.dumps({'action': 'clean-sync-wp16', 'clearAll': True, 'items': items_to_sync})
    try:
        resp = requests.post(GAS_URL, data=payload, headers={'Content-Type': 'text/plain;charset=utf-8'}, timeout=90)
        print(f"📡 ผลตอบกลับจาก Google Apps Script (HTTP {resp.status_code}):")
        print(f"   {resp.text[:400]}")
    except Exception as e:
        print(f"❌ เกิดข้อผิดพลาดในการเชื่อมต่อ GAS: {e}")

    print("\n🎉 การนำเข้าข้อมูลลงใน WP16_งานค้าง เสร็จสมบูรณ์เรียบร้อยแล้ว!")
    return True

if __name__ == '__main__':
    run_import()
